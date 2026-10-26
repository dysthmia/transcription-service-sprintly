from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pymupdf
import pytest
from docx import Document

from app.core.exceptions import ContentError
from app.features.extraction.base import DetectedType
from app.features.extraction.detector import detect_type
from app.features.extraction.service import ExtractionService


class FakeVisionClient:
    def __init__(self, response: str = "Распознанный текст") -> None:
        self.response = response
        self.calls: list[tuple[bytes, str]] = []

    async def extract_vision_markdown(self, image: bytes, mime_type: str) -> str:
        self.calls.append((image, mime_type))
        return self.response


@pytest.fixture
def settings() -> SimpleNamespace:
    return SimpleNamespace(max_document_file_size_mb=1, max_pdf_pages=3)


@pytest.fixture
def service(
    settings: SimpleNamespace, monkeypatch: pytest.MonkeyPatch
) -> ExtractionService:
    mime_types = {
        ".txt": DetectedType.TXT,
        ".docx": DetectedType.DOCX,
        ".pdf": DetectedType.PDF,
    }
    monkeypatch.setattr(
        "app.features.extraction.service.detect_type",
        lambda source: mime_types[source.suffix],
    )
    return ExtractionService(settings, FakeVisionClient())


@pytest.mark.asyncio
async def test_extracts_utf8_and_cp1251_txt(service: ExtractionService, tmp_path: Path) -> None:
    utf8_source = tmp_path / "source.txt"
    utf8_source.write_text("Привет, мир", encoding="utf-8-sig")
    cp1251_source = tmp_path / "legacy.txt"
    cp1251_source.write_bytes("Текст на русском".encode("cp1251"))

    assert (await service.extract(utf8_source)).content == "Привет, мир"
    assert (await service.extract(cp1251_source)).content == "Текст на русском"


@pytest.mark.asyncio
async def test_extracts_docx_paragraphs_and_markdown_table(
    service: ExtractionService,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Первый абзац")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Понятие"
    table.cell(0, 1).text = "Значение"
    table.cell(1, 0).text = "Скорость"
    table.cell(1, 1).text = "42"
    document.add_paragraph("Последний абзац")
    document.save(source)

    result = await service.extract(source)

    assert result.content.index("Первый абзац") < result.content.index("| Понятие | Значение |")
    assert result.content.index("| Понятие | Значение |") < result.content.index("Последний абзац")
    assert "| Скорость | 42 |" in result.content
    assert result.metadata == {"table_count": 1}


@pytest.mark.asyncio
async def test_extracts_pdf_text_layer(service: ExtractionService, tmp_path: Path) -> None:
    source = tmp_path / "source.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "PDF text layer example")
    document.save(source)
    document.close()

    result = await service.extract(source)

    assert "PDF text layer example" in result.content
    assert result.detected_type is DetectedType.PDF
    assert not result.metadata["needs_vision_fallback"]


@pytest.mark.asyncio
async def test_empty_txt_raises_empty_content(service: ExtractionService, tmp_path: Path) -> None:
    source = tmp_path / "empty.txt"
    source.write_text(" \n\t", encoding="utf-8")

    with pytest.raises(ContentError, match="empty_content"):
        await service.extract(source)


@pytest.mark.asyncio
async def test_scanned_pdf_uses_vision_for_only_empty_text_pages(
    settings: SimpleNamespace,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "scanned.pdf"
    document = pymupdf.open()
    document.new_page()
    page = document.new_page()
    page.insert_text((72, 72), "Text layer must not use vision")
    document.save(source)
    document.close()
    monkeypatch.setattr(
        "app.features.extraction.service.detect_type", lambda _: DetectedType.PDF
    )
    vision = FakeVisionClient("| Понятие | Значение |\n| --- | --- |\n| $x$ | 1 |")
    extraction_service = ExtractionService(settings, vision)

    result = await extraction_service.extract(source)

    assert "Text layer must not use vision" in result.content
    assert "| Понятие | Значение |" in result.content
    assert result.metadata["needs_vision_fallback"]
    assert len(vision.calls) == 1
    assert vision.calls[0][1] == "image/png"


@pytest.mark.asyncio
async def test_document_size_and_pdf_page_limits_raise_file_too_large(
    service: ExtractionService,
    settings: SimpleNamespace,
    tmp_path: Path,
) -> None:
    oversized = tmp_path / "large.txt"
    oversized.write_bytes(b"x")
    settings.max_document_file_size_mb = 0

    with pytest.raises(ContentError, match="file_too_large"):
        await service.extract(oversized)

    settings.max_document_file_size_mb = 1
    settings.max_pdf_pages = 1
    pdf_source = tmp_path / "many-pages.pdf"
    document = pymupdf.open()
    document.new_page()
    document.new_page()
    document.save(pdf_source)
    document.close()

    with pytest.raises(ContentError, match="file_too_large"):
        await service.extract(pdf_source)


@pytest.mark.asyncio
@pytest.mark.parametrize("suffix, detected_type", [(".png", DetectedType.PNG), (".jpg", DetectedType.JPEG)])
async def test_image_uses_vision_and_keeps_markdown(
    settings: SimpleNamespace,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    suffix: str,
    detected_type: DetectedType,
) -> None:
    source = tmp_path / f"source{suffix}"
    source.write_bytes(b"image-data")
    monkeypatch.setattr("app.features.extraction.service.detect_type", lambda _: detected_type)
    vision = FakeVisionClient("| Формула | Значение |\n| --- | --- |\n| $E=mc^2$ | энергия |")

    result = await ExtractionService(settings, vision).extract(source)

    assert result.content == vision.response
    assert result.detected_type is detected_type
    assert vision.calls == [(b"image-data", "image/png" if suffix == ".png" else "image/jpeg")]


def test_detector_rejects_unsupported_mime(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "unknown.bin"
    source.write_bytes(b"unknown")
    monkeypatch.setattr(
        "app.features.extraction.detector._mime_type",
        lambda _: "application/octet-stream",
    )

    with pytest.raises(ContentError, match="unsupported_format"):
        detect_type(source)
